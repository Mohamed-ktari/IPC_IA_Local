# docx_parser.py
# Two distinct read modes for .docx files:
#
#   parse_for_ingestion()   — flattens a .docx into the same {"full_text",
#                             "total_pages", "word_count"} shape PDFParser
#                             returns, so chunker.py/ingestion.py work
#                             unmodified. Used for past mémoires/programmes.
#
#   parse_structure_docx()  — reads an rc_agent-generated (or user-edited)
#                             structure skeleton into an ordered list of
#                             {level, title, existing_content} nodes, one
#                             per heading. Used by generation_agent to know
#                             what to fill.
#
# Both share the same heading-detection primitive: docx_writer.py's
# comment confirms the template's Titre1..Titre9 styles are exposed by
# python-docx as "Heading 1".."Heading 9" — so we detect level purely from
# paragraph.style.name, never from indentation or numbering (a user
# hand-editing the file can't be trusted to keep either consistent).

import re
from pathlib import Path

from docx import Document

HEADING_STYLE_RE = re.compile(r"^Heading (\d+)$")


def _heading_level(paragraph) -> int | None:
    # Returns the heading level (1-9) if this paragraph is a heading,
    # else None. Falls back to checking for "Titre" in the style name
    # in case a future template swap changes how python-docx exposes it —
    # belt-and-suspenders per your request to not assume the template
    # never changes.
    style_name = paragraph.style.name if paragraph.style else ""
    m = HEADING_STYLE_RE.match(style_name)
    if m:
        return int(m.group(1))

    m2 = re.match(r"^Titre\s*(\d+)$", style_name, re.IGNORECASE)
    if m2:
        return int(m2.group(1))

    return None


def parse_for_ingestion(file_path: str | Path) -> dict:
    """
    Flattens a .docx into markdown-style text so it can be chunked and
    embedded the same way PDFParser output is. Used for ingesting past
    mémoires and programmes into Chroma.
    """
    file_path = Path(file_path)
    doc = Document(file_path)

    lines = []
    word_count = 0

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue

        level = _heading_level(p)
        word_count += len(text.split())

        if level is not None:
            lines.append(f"{'#' * min(level, 6)} {text}")
        else:
            lines.append(text)

    full_text = "\n\n".join(lines)

    return {
        "full_text": full_text,
        # .docx has no fixed page concept the way a PDF does — leaving this
        # as 1 rather than None keeps downstream code that expects an int
        # (metadata["total_pages"]) from breaking. Revisit if page-range
        # slicing ever needs to apply to docx sources too.
        "total_pages": 1,
        "word_count": word_count,
    }


def parse_structure_docx(file_path: str | Path) -> list[dict]:
    """
    Reads a structure .docx (built by docx_writer.generate_memoire_docx,
    possibly hand-edited by the user afterward) into an ordered list of:
        {"level": int, "title": str, "existing_content": str}

    existing_content is whatever text sits between this heading and the
    next one — normally the blank paragraph docx_writer inserted, but if
    the user typed a free-text instruction ("mets l'accent sur X") or
    already wrote real content there, it's captured as-is. Downstream
    code decides what to do with it (instruction vs already-filled).

    Resilient to a user having deleted/reordered/added headings — this
    just walks paragraphs in document order and doesn't assume any fixed
    section list.
    """
    file_path = Path(file_path)
    doc = Document(file_path)

    nodes: list[dict] = []
    current: dict | None = None

    for p in doc.paragraphs:
        level = _heading_level(p)
        text = p.text.strip()

        if level is not None:
            if current is not None:
                current["existing_content"] = current["existing_content"].strip()
                nodes.append(current)
            current = {"level": level, "title": text, "existing_content": ""}
        else:
            if current is None:
                # Text before the first heading (e.g. template's title/date/
                # logo paragraphs) — not a section, ignore for this purpose.
                continue
            if text:
                current["existing_content"] += (text + "\n")

    if current is not None:
        current["existing_content"] = current["existing_content"].strip()
        nodes.append(current)

    return nodes