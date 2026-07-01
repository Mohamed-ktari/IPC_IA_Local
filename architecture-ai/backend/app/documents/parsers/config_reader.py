"""
config_reader.py
----------------
Reads the human-provided Excel configuration file.

Sheet layout ("Configuration"):

  Champ                | Valeur | Description
  Page début            | 8      | ...
  Page fin               | 13     | ...

  (blank row)

  Colonne source         | Colonne normalisée
  Localisation             | localisation
  n de sondage             | reference_echantillon
  Description               | description_materiau
  Echantillon                | echantillon_flag
  N Echantillon               | reference_labo
  Conclusion                   | resultat

The second table is located by its header row ("Colonne source" /
"Colonne normalisée") rather than a fixed row number, so the user can
add/remove mapping rows freely without breaking parsing. Each row below
that header is one source→target pair. Blank rows are skipped. A row
missing either cell is reported and skipped rather than raising — partial
typos in one row shouldn't kill the whole config.
"""

import unicodedata
from pathlib import Path

import openpyxl


_LABEL_MAP = {
    "page debut":      "page_start",
    "page de debut":   "page_start",
    "debut":           "page_start",
    "page start":      "page_start",
    "premiere page":   "page_start",
    "page fin":        "page_end",
    "page de fin":     "page_end",
    "fin":             "page_end",
    "page end":        "page_end",
    "derniere page":   "page_end",
}

_MAPPING_HEADER_SOURCE = {"colonne source", "source column", "colonne", "source"}
_MAPPING_HEADER_TARGET = {"colonne normalisee", "champ normalise", "target field", "normalise"}


def _normalize(text: str) -> str:
    nfkd = unicodedata.normalize("NFKD", str(text).lower().strip())
    return "".join(c for c in nfkd if not unicodedata.combining(c))


class ConfigReader:

    def read(self, file_path: str | Path) -> dict:
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"Config file not found: {file_path}")
        if file_path.suffix.lower() not in (".xlsx", ".xls"):
            raise ValueError(f"Config file must be .xlsx or .xls, got: {file_path.suffix}")

        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb["Configuration"] if "Configuration" in wb.sheetnames else wb.active

        all_rows = list(ws.iter_rows(min_col=1, max_col=2, values_only=True))

        page_fields    = self._read_page_fields(all_rows)
        column_mapping = self._read_mapping_table(all_rows)

        config = {
            "page_start":     page_fields.get("page_start"),
            "page_end":       page_fields.get("page_end"),
            "column_mapping": column_mapping or None,
        }

        if config["page_start"] is not None and config["page_end"] is not None:
            if config["page_start"] > config["page_end"]:
                print(
                    f"[ConfigReader] Warning: page_start ({config['page_start']}) "
                    f"> page_end ({config['page_end']}) — both ignored"
                )
                config["page_start"] = None
                config["page_end"]   = None

        if not config["column_mapping"]:
            print(
                "[ConfigReader] Warning: no valid 'Colonne source / Colonne "
                "normalisée' mapping rows found — extraction will fail without one."
            )

        return config

    # ------------------------------------------------------------------
    # Top key/value table: Page début / Page fin
    # ------------------------------------------------------------------

    def _read_page_fields(self, rows: list[tuple]) -> dict:
        result = {}
        for label_cell, value_cell in rows:
            if label_cell is None:
                continue
            key = _LABEL_MAP.get(_normalize(str(label_cell)))
            if key is None or value_cell is None:
                continue
            try:
                val = int(float(value_cell))
                if val >= 1:
                    result[key] = val
                else:
                    print(f"[ConfigReader] Warning: {key}={val} is < 1, ignored")
            except (ValueError, TypeError):
                print(f"[ConfigReader] Warning: could not parse {key}={value_cell!r}")
        return result

    # ------------------------------------------------------------------
    # Second table: Colonne source | Colonne normalisée
    # ------------------------------------------------------------------

    def _read_mapping_table(self, rows: list[tuple]) -> dict:
        header_idx = None
        for i, (a, b) in enumerate(rows):
            if a is None or b is None:
                continue
            if _normalize(a) in _MAPPING_HEADER_SOURCE and _normalize(b) in _MAPPING_HEADER_TARGET:
                header_idx = i
                break

        if header_idx is None:
            print(
                "[ConfigReader] Warning: mapping table header "
                "('Colonne source' / 'Colonne normalisée') not found."
            )
            return {}

        mapping = {}
        for source_cell, target_cell in rows[header_idx + 1:]:
            if source_cell is None and target_cell is None:
                continue  # blank row — just skip, don't stop (user may have gaps)

            source = str(source_cell).strip() if source_cell is not None else ""
            target = str(target_cell).strip() if target_cell is not None else ""

            if not source and not target:
                continue
            if not source or not target:
                print(
                    f"[ConfigReader] Warning: incomplete mapping row "
                    f"(source={source!r}, target={target!r}) — skipped. "
                    f"Both columns must be filled."
                )
                continue

            normalized_target = self._sanitize_target(target)
            if normalized_target != target:
                print(
                    f"[ConfigReader] Note: target field {target!r} normalised "
                    f"to {normalized_target!r} (spaces/accents not allowed in field names)."
                )

            mapping[source] = normalized_target

        return mapping

    def _sanitize_target(self, target: str) -> str:
        """
        Target field names become JSON keys — keep them safe even if the
        user types accents, spaces, or punctuation (e.g. "Référence Labo"
        → "reference_labo"). Prevents downstream JSON-key issues without
        rejecting the row outright.
        """
        norm = _normalize(target)
        norm = norm.replace(" ", "_").replace("-", "_")
        norm = "".join(c for c in norm if c.isalnum() or c == "_")
        norm = norm.strip("_")
        return norm or target  # fallback to original if sanitizing empties it

    # ------------------------------------------------------------------
    @staticmethod
    def generate_template(output_path: str | Path = "config_template.xlsx") -> Path:
        output_path = Path(output_path)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Configuration"

        from openpyxl.styles import Font, PatternFill, Alignment

        header_font = Font(bold=True, color="FFFFFF", name="Arial", size=11)
        header_fill = PatternFill("solid", fgColor="1F4E79")
        label_font  = Font(bold=True, name="Arial", size=10)
        value_font  = Font(name="Arial", size=10)
        note_font   = Font(italic=True, color="888888", name="Arial", size=9)
        center      = Alignment(horizontal="center", vertical="center")
        left        = Alignment(horizontal="left", vertical="center", wrap_text=True)

        # ── Top table: page range ──────────────────────────────────────
        ws.append(["Champ", "Valeur", "Description"])
        for col in range(1, 4):
            cell = ws.cell(row=1, column=col)
            cell.font, cell.fill, cell.alignment = header_font, header_fill, center

        top_rows = [
            ("Page début", "", "Numéro de la première page contenant les matériaux (ex: 8)"),
            ("Page fin",   "", "Numéro de la dernière page contenant les matériaux (ex: 13)"),
        ]
        r = 2
        for label, value, description in top_rows:
            ws.cell(row=r, column=1, value=label).font = label_font
            ws.cell(row=r, column=1).alignment = left
            ws.cell(row=r, column=2, value=value).font = value_font
            ws.cell(row=r, column=2).alignment = left
            ws.cell(row=r, column=3, value=description).font = note_font
            ws.cell(row=r, column=3).alignment = left
            r += 1

        r += 1  # blank separator row

        # ── Second table: column mapping ─────────────────────────────────
        ws.cell(row=r, column=1, value="Colonne source").font = header_font
        ws.cell(row=r, column=1).fill = header_fill
        ws.cell(row=r, column=1).alignment = center
        ws.cell(row=r, column=2, value="Colonne normalisée").font = header_font
        ws.cell(row=r, column=2).fill = header_fill
        ws.cell(row=r, column=2).alignment = center
        ws.cell(row=r, column=3, value="Description").font = header_font
        ws.cell(row=r, column=3).fill = header_fill
        ws.cell(row=r, column=3).alignment = center
        r += 1

        example_rows = [
            ("Localisation", "localisation",
             "colonne source est le nom de la colonne TEL QU'ELLE APPARAÎT dans le document source"),
            ("Description", "description_materiau",
             "colonne normalisée est le nom du champ dans le résultat final (sans accents/espaces de préférence)"),
            ("Conclusion", "resultat",
             "Utiliser exactement 'resultat' si possible pour le calcul automatique présence/absence"),
            ("", "", "Ajoutez une ligne par colonne à extraire — autant de lignes que nécessaire"),
        ]
        for source, target, description in example_rows:
            ws.cell(row=r, column=1, value=source).font = value_font
            ws.cell(row=r, column=1).alignment = left
            ws.cell(row=r, column=2, value=target).font = value_font
            ws.cell(row=r, column=2).alignment = left
            ws.cell(row=r, column=3, value=description).font = note_font
            ws.cell(row=r, column=3).alignment = left
            if (r % 2) == 0:
                fill = PatternFill("solid", fgColor="F2F7FC")
                for col in range(1, 4):
                    ws.cell(row=r, column=col).fill = fill
            r += 1

        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 28
        ws.column_dimensions["C"].width = 60
        ws.row_dimensions[1].height = 20

        wb.save(output_path)
        return output_path